from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "doc" / "report" / "รายงานวิชาการ_CourseHub_ฉบับร่าง.docx"
FONT = "Angsana New"
NAVY = RGBColor(23, 43, 69)
GRAY = RGBColor(87, 96, 107)


def font(run, size=16, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + key), FONT)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_border(cell, color="D9DEE5"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for m, val in {"top":top,"start":start,"bottom":bottom,"end":end}.items():
        el = mar.find(qn("w:" + m))
        if el is None:
            el = OxmlElement("w:" + m)
            mar.append(el)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.5)
sec.bottom_margin = Cm(2.3)
sec.left_margin = Cm(3.0)
sec.right_margin = Cm(2.3)
sec.header_distance = Cm(1.2)
sec.footer_distance = Cm(1.2)
sec.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(16)
normal.paragraph_format.line_spacing = 1.16
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.first_line_indent = Cm(0.7)
normal.paragraph_format.widow_control = True

for name, size, before, after in [("Title", 24, 0, 12), ("Heading 1", 20, 15, 8),
                                  ("Heading 2", 18, 10, 5), ("Heading 3", 16, 8, 3)]:
    s = styles[name]
    s.font.name = FONT
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = NAVY if name != "Title" else RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.first_line_indent = Cm(0)

# The built-in Word Title style can carry a theme-colored bottom border.
title_ppr = styles["Title"]._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

for name in ("Caption", "Reference", "Placeholder"):
    if name not in styles:
        styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
styles["Caption"].font.name = FONT
styles["Caption"].font.size = Pt(14)
styles["Caption"].font.color.rgb = GRAY
styles["Caption"].paragraph_format.space_before = Pt(5)
styles["Caption"].paragraph_format.space_after = Pt(10)
styles["Caption"].paragraph_format.first_line_indent = Cm(0)
styles["Reference"].font.name = FONT
styles["Reference"].font.size = Pt(14)
styles["Reference"].paragraph_format.left_indent = Cm(0.7)
styles["Reference"].paragraph_format.first_line_indent = Cm(-0.7)
styles["Reference"].paragraph_format.space_after = Pt(7)
styles["Placeholder"].font.name = FONT
styles["Placeholder"].font.size = Pt(14)
styles["Placeholder"].font.color.rgb = GRAY
styles["Placeholder"].paragraph_format.left_indent = Cm(0.5)
styles["Placeholder"].paragraph_format.right_indent = Cm(0.5)
styles["Placeholder"].paragraph_format.first_line_indent = Cm(0)
styles["Placeholder"].paragraph_format.space_before = Pt(8)
styles["Placeholder"].paragraph_format.space_after = Pt(4)


def para(text="", style=None, align=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), bold=True)
        font(p.add_run(text[len(bold_lead):]))
    else:
        font(p.add_run(text), 14 if style in ("Caption", "Reference", "Placeholder") else 16)
    return p


def heading(text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        if widths:
            c.width = Cm(widths[i])
        c.text = h
        set_cell_shading(c, "E8EEF4")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(c)
        set_cell_margin(c)
        for p in c.paragraphs:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                font(r, 14, True, NAVY)
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            if widths:
                c.width = Cm(widths[i])
            c.text = str(val)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j % 2:
                set_cell_shading(c, "F7F9FB")
            set_cell_border(c)
            set_cell_margin(c)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    font(r, 13)
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(header_flag)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def figure(number, description, instruction):
    para(f"[ตำแหน่งรูปที่ {number}  {instruction}]", "Placeholder")
    para(f"รูปที่ {number} {description}", "Caption", WD_ALIGN_PARAGRAPH.CENTER)


def page_break():
    doc.add_page_break()


def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "คลิกขวาเพื่ออัปเดตสารบัญ"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, txt, end])


# Cover
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("รายงานวิชาการโครงงาน\nการออกแบบและพัฒนาระบบแนะนำคอร์สเรียน\nคอร์สดีบอกต่อ"), 24, True)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Design and Development of CourseHub Course Recommendation System"), 17)
doc.add_paragraph("\n\n")
para("จัดทำโดย กลุ่ม 36", align=WD_ALIGN_PARAGRAPH.CENTER)
for line in ["673380054-1  นายภาคิน เมฆสุวรรณ", "673380072-9  นายเกียรติศักดิ์ นันทรัตน์",
             "673380062-2  นายศุภวัฒน์ ข่ายทอง", "673380064-8  นายสรวิชญ์ วันเสน"]:
    para(line, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph("\n")
para("รายงานประกอบรายวิชา CP353002 Principles of Software Design and Development", align=WD_ALIGN_PARAGRAPH.CENTER)
para("สาขาวิชาวิทยาการคอมพิวเตอร์ วิทยาลัยการคอมพิวเตอร์ มหาวิทยาลัยขอนแก่น", align=WD_ALIGN_PARAGRAPH.CENTER)
para("[เติมภาคเรียน ปีการศึกษา ชื่ออาจารย์ผู้สอน และวันที่ส่งก่อนยื่นรายงาน]", "Placeholder", WD_ALIGN_PARAGRAPH.CENTER)
para("ฉบับร่างสำหรับตรวจทานข้อมูล ณ 8 ตุลาคม พ.ศ. 2569", align=WD_ALIGN_PARAGRAPH.CENTER)
page_break()

heading("บทคัดย่อ")
para("รายงานนี้นำเสนอการออกแบบ พัฒนา และตรวจสอบระบบคอร์สดีบอกต่อ (CourseHub) ซึ่งเป็นเว็บสำหรับค้นหาและคัดเลือกคอร์สเรียนจากผู้ให้บริการหลายราย ปัญหาที่ศึกษาเป็นการกระจายตัวของข้อมูลคอร์สและความยากในการเลือกคอร์สให้ตรงกับหมวดหมู่ ระดับ ภาษา งบประมาณ และเวลาเรียนที่ผู้เรียนมี ระบบต้นแบบประกอบด้วยบัญชีผู้ใช้ แค็ตตาล็อก การบันทึกคอร์ส การจัดการผู้ให้บริการและคอร์ส การตรวจอนุมัติ และตัวจับคู่คอร์สแบบอิงกฎ โดยการชำระเงินและเนื้อหาการเรียนยังอยู่ที่เว็บไซต์ต้นทาง")
para("วิธีดำเนินงานใช้การวิเคราะห์ข้อกำหนดจากแผนและเอกสารโครงการ ตรวจสอบโค้ดและฐานข้อมูล แล้วสังเคราะห์หลักการระบบแนะนำ การออกแบบซอฟต์แวร์ และเกณฑ์คุณภาพที่เกี่ยวข้อง สถาปัตยกรรมใช้ React และ TypeScript สำหรับส่วนติดต่อผู้ใช้ Spring Boot สำหรับ REST API และกฎธุรกิจ และ PostgreSQL สำหรับข้อมูลสัมพันธ์ ตัวจับคู่กรองคอร์สที่ไม่เข้าเงื่อนไขก่อน แล้วเฉลี่ยคะแนนด้านราคา เวลาเรียน และคุณภาพรีวิวที่เผยแพร่ พร้อมแสดงเหตุผลประกอบคำแนะนำ")
para("หลักฐานในโปรเจกต์ ณ วันที่ตรวจสอบมีรายงานผล backend 308 กรณีผ่านทั้งหมด และผล Playwright 10 กรณีผ่านโดยไม่พบ skipped หรือ flaky ในรอบล่าสุดที่ตรวจได้ ผลดังกล่าวสนับสนุนความถูกต้องของพฤติกรรมที่ทดสอบ แต่ยังไม่พิสูจน์ความเหมาะสมของคำแนะนำต่อผู้เรียนจริง เนื่องจากยังไม่มีชุดข้อมูลประเมินความเกี่ยวข้องโดยมนุษย์หรือการทดลองกับผู้ใช้ ระบบยังมีข้อจำกัดด้านคอร์สแบบสมาชิก สกุลเงินอื่น การประเมินบนอุปกรณ์หลากหลาย และการยืนยันการใช้งานส่วนหน้าในระบบสาธารณะ จึงเสนอให้เก็บข้อมูลและประเมินผลกับผู้เรียนก่อนอ้างคุณภาพเชิงการแนะนำ")
para("คำสำคัญ: ระบบแนะนำคอร์สเรียน, การจับคู่ตามความต้องการ, การออกแบบซอฟต์แวร์, การทดสอบระบบ, CourseHub")
heading("Abstract")
para("This report presents the design, implementation, and verification of CourseHub, a web application for discovering courses from multiple providers. The prototype supports a course catalog, account and provider management, moderation, bookmarks, and a rule-based course matcher. The matcher first applies hard constraints on category, level, language, and budget, then ranks eligible courses using price fit, study effort, and published-review quality. It returns up to three courses with explicit reasons. The system uses React and TypeScript, a Spring Boot REST API, and PostgreSQL. Repository evidence inspected for this report records 308 passing backend tests and a recent Playwright run with 10 passing scenarios and no skipped or flaky tests. These results establish tested functional behavior, but they do not establish recommendation relevance for real learners. Human relevance judgments, usability assessment, broader browser coverage, and verified public frontend deployment remain future work.")
para("Keywords: course recommendation, rule-based matching, software design, software testing, CourseHub")
page_break()

heading("คำนำ")
para("รายงานวิชาการฉบับนี้จัดทำขึ้นเพื่อเสนอแนวคิดและกระบวนการพัฒนาระบบคอร์สดีบอกต่อ (CourseHub) ซึ่งมุ่งช่วยให้ผู้เรียนค้นหาและพิจารณาคอร์สเรียนที่เหมาะสมกับความสนใจ ระดับความรู้ ภาษา และงบประมาณของตน เนื้อหาครอบคลุมที่มาและความสำคัญของปัญหา แนวคิดและเทคโนโลยีที่เกี่ยวข้อง การวิเคราะห์ความต้องการ การออกแบบและพัฒนาระบบ ตลอดจนผลการทดสอบฟังก์ชันที่ดำเนินการแล้ว")
para("การเรียบเรียงรายงานอาศัยแผนการดำเนินงาน เอกสารประกอบโครงการ ซอร์สโค้ด และหลักฐานการทดสอบ เพื่อแสดงความเชื่อมโยงระหว่างเป้าหมาย วิธีการพัฒนา และผลที่ตรวจสอบได้ ผู้จัดทำได้ระบุข้อจำกัดของระบบและแนวทางพัฒนาต่อไว้ด้วย เพื่อให้รายงานสะท้อนสถานะของโครงการตามความเป็นจริง และเป็นประโยชน์ต่อการศึกษาและพัฒนาระบบแนะนำคอร์สเรียนในอนาคต")
para("ผู้จัดทำ\nกลุ่ม 36")
page_break()

heading("สารบัญ")
p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0)
field(p, ' TOC \\o "1-3" \\h \\z \\u ')
heading("สารบัญภาพ")
for x in ["รูปที่ 3.1 ภาพรวมองค์ประกอบของระบบ CourseHub", "รูปที่ 3.2 แผนภาพความสัมพันธ์ข้อมูลหลัก",
          "รูปที่ 4.1 หน้าค้นหาคอร์สและตัวกรอง", "รูปที่ 4.2 หน้าคำถามและผลลัพธ์ตัวจับคู่คอร์ส",
          "รูปที่ 4.3 หน้าจัดการคอร์สและการตรวจอนุมัติ", "รูปที่ 4.4 หลักฐานผลการทดสอบและ CI"]:
    para(x)
page_break()

heading("บทที่ 1 บทนำ")
heading("1.1 ที่มาและความสำคัญ", 2)
para("ผู้เรียนที่ต้องการพัฒนาทักษะมักพบคอร์สจากหลายแพลตฟอร์มและหลายผู้ให้บริการ โดยแต่ละแห่งแสดงรายละเอียด ราคา ภาษา และรูปแบบการเรียนต่างกัน การค้นหาด้วยคำสำคัญเพียงอย่างเดียวจึงอาจไม่เพียงพอสำหรับการตัดสินใจเมื่อผู้เรียนมีข้อจำกัดด้านงบประมาณและเวลา แผนโครงการ CourseHub จึงกำหนดระบบตัวกลางที่รวบรวมข้อมูลคอร์ส ให้ผู้ให้บริการจัดการข้อมูลผ่านการตรวจอนุมัติ และช่วยผู้เรียนคัดเลือกคอร์สจากความต้องการที่ระบุ [1–3]")
para("ระบบนี้ทำหน้าที่ค้นหา เปรียบเทียบ บันทึก และเชื่อมออกไปยังต้นทางคอร์ส ไม่ได้รับชำระเงินหรือจัดการเนื้อหาการเรียนในตัวเอง ขอบเขตดังกล่าวช่วยให้แบบจำลองข้อมูลและสิทธิ์การใช้งานสอดคล้องกับสิ่งที่พัฒนาได้จริง ขณะเดียวกันข้อมูลที่เผยแพร่ต้องผ่านกระบวนการตรวจของผู้ดูแล เพื่อป้องกันการแสดงคอร์สที่ยังไม่พร้อมหรือผู้ให้บริการที่ยังไม่รับรอง [1,3]")
heading("1.2 วัตถุประสงค์", 2)
for s in [
    "1. ออกแบบและพัฒนาเว็บสำหรับค้นหา กรอง บันทึก และดูรายละเอียดคอร์สจากผู้ให้บริการหลายราย",
    "2. พัฒนากลไกจัดการผู้ให้บริการ คอร์ส รีวิว และการตรวจอนุมัติตามบทบาทผู้ใช้",
    "3. พัฒนาตัวจับคู่คอร์สแบบอิงกฎที่แสดงคะแนนและเหตุผลของคำแนะนำ",
    "4. ตรวจสอบพฤติกรรมของระบบด้วยการทดสอบระดับหน่วย การเชื่อมต่อ และการใช้งานตั้งแต่ต้นจนจบ",
    "5. วิเคราะห์ข้อจำกัดของหลักฐานและกำหนดแนวทางประเมินกับผู้เรียนจริงในอนาคต",
]: para(s)
heading("1.3 ขอบเขตของโครงงาน", 2)
para("ผู้ใช้แบ่งเป็นผู้เรียน สมาชิกทีมผู้ให้บริการ และผู้ดูแลระบบ ผู้เรียนค้นหา กรองและบันทึกคอร์ส ใช้แบบสอบถามตัวจับคู่ และส่งรีวิวผ่าน API ได้ ผู้ให้บริการสร้างและแก้ไขข้อมูลสถาบันและคอร์สดราฟต์ก่อนส่งตรวจ ผู้ดูแลรับรองผู้ให้บริการ ตรวจคอร์สและรีวิว ระบบแสดงเฉพาะคอร์สที่เผยแพร่จากผู้ให้บริการที่มีสถานะใช้งาน [2,3]")
para("ขอบเขตของตัวจับคู่ใช้ข้อมูลหมวดหมู่ ระดับ ภาษา งบประมาณ และชั่วโมงเรียนต่อสัปดาห์ แล้วจัดอันดับจากราคา เวลาเรียน และคะแนนรีวิวที่เผยแพร่ ไม่มีการฝึกโมเดลปัญญาประดิษฐ์หรือใช้ประวัติพฤติกรรมจำนวนมากเพื่อเรียนรู้ความชอบโดยอัตโนมัติ ฟังก์ชัน Roadmap และหน้าจอผู้เรียนสำหรับเขียนรีวิวยังไม่อยู่ในสถานะสมบูรณ์ตามหลักฐานที่ตรวจ [3,7]")
heading("1.4 คำถามในการศึกษา", 2)
para("การศึกษาพิจารณาว่าโครงสร้างระบบสามารถแยกสิทธิ์และวงจรชีวิตข้อมูลคอร์สได้เพียงใด ตัวจับคู่ให้ผลที่สอดคล้องกับข้อจำกัดที่ผู้ใช้ระบุหรือไม่ และหลักฐานทดสอบยืนยันพฤติกรรมใดได้บ้าง คำถามเกี่ยวกับความพึงพอใจหรือความแม่นยำเชิงความเกี่ยวข้องของคำแนะนำยังต้องอาศัยการศึกษาในผู้ใช้จริง")
heading("1.5 ประโยชน์ที่คาดว่าจะได้รับ", 2)
para("โครงงานให้ต้นแบบที่ผู้เรียนใช้สำรวจคอร์สตามเงื่อนไขของตน และเป็นตัวอย่างของการออกแบบ REST API ฐานข้อมูลสัมพันธ์ การจัดการสิทธิ์ รูปแบบการออกแบบ และการทดสอบระบบอย่างตรวจสอบย้อนกลับได้ ประโยชน์เชิงผลลัพธ์ต่อการตัดสินใจเรียนยังเป็นสมมติฐานที่ต้องประเมินต่อ")
page_break()

heading("บทที่ 2 แนวคิดและงานที่เกี่ยวข้อง")
heading("2.1 ระบบแนะนำและการจับคู่ตามข้อจำกัด", 2)
para("งานสำรวจของ Adomavicius และ Tuzhilin [9] อธิบายแนวทางระบบแนะนำหลายกลุ่ม รวมถึงการใช้เนื้อหาและความสัมพันธ์ระหว่างผู้ใช้กับรายการ ระบบ CourseHub เลือกแนวทางที่อาศัยข้อมูลคอร์สและความต้องการที่ผู้เรียนระบุเอง เพราะมีข้อมูลรีวิวและพฤติกรรมผู้ใช้จำกัดในระยะต้น การเลือกนี้เป็นเหตุผลด้านความพร้อมของข้อมูลและความตรวจสอบได้ ไม่ใช่หลักฐานว่าคะแนนแนะนำดีกว่าแนวทางอื่น")
para("Burke [10] ชี้ว่าการผสมวิธีแนะนำมีหลายรูปแบบและควรเลือกตามปัญหาและข้อมูลที่มี CourseHub ยังไม่ได้ใช้ collaborative filtering หรือระบบผสมแบบงานดังกล่าว กลไกที่พัฒนาเป็นการกรองตามข้อจำกัดก่อนคำนวณคะแนน ซึ่งช่วยให้คอร์สที่ขัดเงื่อนไขสำคัญไม่หลุดเข้าสู่ผลลัพธ์")
heading("2.2 การอธิบายเหตุผลของคำแนะนำ", 2)
para("Tintarev และ Masthoff [11] อภิปรายบทบาทของคำอธิบายในระบบแนะนำ เช่น การช่วยให้ผู้ใช้เข้าใจและประเมินข้อเสนอ ระบบ CourseHub จึงคืนเหตุผลตามแต่ละองค์ประกอบคะแนน ได้แก่ ราคา เวลาเรียน และรีวิว รวมทั้งเหตุผลเมื่อไม่พบคอร์ส อย่างไรก็ตาม รายงานนี้ยังไม่มีผลทดลองว่าข้อความเหล่านี้ช่วยให้ผู้เรียนตัดสินใจได้ดีขึ้นจริง")
heading("2.3 การประเมินคุณภาพซอฟต์แวร์", 2)
para("กรอบ ISO/IEC 25010:2023 [12] ใช้คุณลักษณะของผลิตภัณฑ์ซอฟต์แวร์เป็นแนวทางกำหนดการประเมิน รายงานนี้ตรวจหลักฐานด้านความเหมาะสมเชิงหน้าที่ ความปลอดภัย และความสามารถในการบำรุงรักษาผ่านโค้ดและการทดสอบเป็นหลัก การผ่านเทสต์ไม่ใช่การประเมินครบทุกคุณลักษณะ โดยเฉพาะประสบการณ์ผู้ใช้ ประสิทธิภาพภายใต้ภาระงาน และความพร้อมใช้งานบนคลาวด์")
para("Shani และ Gunawardana [13] แยกวิธีประเมินระบบแนะนำตามเป้าหมายและบริบท ดังนั้นรายงานนี้ใช้ผลทดสอบเชิงฟังก์ชันเพื่อยืนยันการทำงานตามกฎ แต่ยังไม่รายงาน precision, recall หรือคุณภาพอันดับ เพราะไม่มีชุดคำตอบอ้างอิงจากผู้เรียน การใส่ตัวเลขเหล่านี้โดยไม่มีข้อมูลจริงจะทำให้ข้อสรุปเกินหลักฐาน")
heading("2.4 หลักการออกแบบระบบบริการ", 2)
para("สถาปัตยกรรมแบบ Service Layer และ Data Transfer Object ตาม Fowler [14] ช่วยแยกขอบเขตการประมวลผลจากการส่งข้อมูลผ่านเครือข่าย ใน CourseHub ส่วนควบคุม HTTP เรียกบริการที่บังคับกฎธุรกิจและใช้ repository เข้าถึงข้อมูล การแยกนี้ช่วยให้ทดสอบและเปลี่ยนรายละเอียดภายในได้โดยลดผลกระทบต่อสัญญา API [4]")
heading("2.5 สังเคราะห์แนวคิดเพื่อการออกแบบ", 2)
table(["แหล่งศึกษา", "ประเด็นที่นำมาใช้", "การประยุกต์ใน CourseHub", "สิ่งที่ยังไม่พิสูจน์"], [
    ("Adomavicius และ Tuzhilin [9]", "ประเภทและข้อจำกัดของระบบแนะนำ", "จับคู่จากข้อมูลคอร์สและความต้องการที่ระบุ", "คุณภาพเทียบวิธีอื่น"),
    ("Burke [10]", "การเลือกและผสมวิธีแนะนำ", "แยกขั้นกรองกับขั้นจัดอันดับ", "ยังไม่มีระบบผสมหรือการทดลองเปรียบเทียบ"),
    ("Tintarev และ Masthoff [11]", "คำอธิบายผลแนะนำ", "แสดงเหตุผลราคา เวลา และรีวิว", "ความเข้าใจของผู้ใช้จริง"),
    ("ISO/IEC 25010 [12]", "กรอบคุณภาพผลิตภัณฑ์", "ทดสอบหน้าที่ สิทธิ์ และความคงทนของข้อมูล", "คุณภาพทุกมิติในสภาพใช้งานจริง"),
], [3.1, 3.5, 4.2, 4.0])
page_break()

heading("บทที่ 3 วิธีดำเนินงานและการออกแบบระบบ")
heading("3.1 วิธีการศึกษาและแหล่งข้อมูล", 2)
para("การศึกษาใช้แนวทางออกแบบและพัฒนาระบบ โดยเริ่มจากหน้าแผน Notion ที่ผู้จัดทำส่งเป็น PDF ซึ่งปรับแผนวันที่ 15 กันยายน พ.ศ. 2569 และตารางข้อกำหนด จากนั้นตรวจโค้ด เอกสารสถาปัตยกรรม สคริปต์ฐานข้อมูล แผนทดสอบ และรายงานผลทดสอบใน repository ณ วันที่ 8 ตุลาคม พ.ศ. 2569 ข้อความในแผนเป็นเป้าหมาย ณ วันที่จัดทำ ไม่ใช่หลักฐานว่าฟังก์ชันเสร็จแล้ว ข้อสรุปว่า ‘พัฒนาแล้ว’ หรือ ‘ทดสอบผ่าน’ จึงต้องมีโค้ดหรือผลทดสอบรองรับ [1–8]")
heading("3.1.1 แผนดำเนินงานที่ใช้กำกับโครงการ", 3)
para("แผน Notion แบ่งงานเป็น P0 ถึง P6 ในกรอบตัวอย่างแปดสัปดาห์ โดยไม่ได้ระบุวันส่งจริง ตารางนี้สรุปลำดับและผลลัพธ์ที่แผนคาดหวัง จึงไม่ควรอ่านเป็นบันทึกวันที่ทำงานจริง [1]")
table(["ช่วง", "งานตามแผน", "หลักฐานที่ควรตรวจเมื่อปิดช่วง"], [
    ("P0", "ยืนยันขอบเขต ทีม และข้อกำหนด", "README, requirement matrix, use cases"),
    ("P1", "วางรากฐาน Spring Boot, React, schema และ security", "build, migration, auth และ deploy skeleton"),
    ("P2", "Provider และ Course CRUD, catalog, bookmark", "API/UI และการตรวจ ownership"),
    ("P3", "moderation, review, audit และ observer", "state transitions, rollback และ concurrency tests"),
    ("P4", "matcher, E2E, SOLID และ diagrams", "ผลทดสอบและเอกสารที่อ้างโค้ดจริง"),
    ("P5", "deploy, Swagger และตรวจระบบ", "public URL, smoke test, backup/restore"),
    ("P6", "แก้ข้อบกพร่องและเตรียมส่ง", "รายงาน สไลด์ และ commit ที่ส่งตรงกัน"),
], [1.4, 5.9, 6.8])
p = heading("3.2 ความต้องการหลักและผู้เกี่ยวข้อง", 2)
p.paragraph_format.page_break_before = True
table(["ผู้ใช้", "งานหลัก", "ข้อกำหนดสำคัญ"], [
    ("ผู้เรียน", "ค้นหา กรอง บันทึกคอร์ส ใช้ตัวจับคู่ และรีวิว", "เห็นเฉพาะคอร์สเผยแพร่และข้อมูลของตน"),
    ("ผู้ให้บริการ", "จัดการโปรไฟล์ ทีม และคอร์สดราฟต์", "แก้ไขเฉพาะข้อมูลที่มีสิทธิ์และส่งคอร์สเข้าตรวจ"),
    ("ผู้ดูแล", "รับรองผู้ให้บริการ ตรวจคอร์สและรีวิว", "บันทึกสถานะ เหตุผล และประวัติการกระทำ"),
], [2.5, 5.4, 6.9])
heading("3.3 สถาปัตยกรรมและเทคโนโลยี", 2)
para("ส่วนหน้าพัฒนาเป็น Single Page Application ด้วย React 18, TypeScript และ Vite ส่วนหลังใช้ Java 21 และ Spring Boot 3.5.16 ผ่าน REST API ที่มีคำนำหน้า /api/v1 ข้อมูลเก็บใน PostgreSQL และจัดการ schema ด้วย Flyway การแบ่งชั้นหลักคือ Controller, Service และ Repository โดยใช้ DTO แยกข้อมูล API ออกจาก entity ฐานข้อมูล [2,4,5]")
table(["ชั้น", "หน้าที่", "หลักฐาน"], [
    ("React และ TypeScript", "หน้าค้นหา บัญชี Provider Admin และ Matcher", "code/frontend/src/pages/"),
    ("Spring MVC และ Security", "รับคำขอ ตรวจ session, CSRF และสิทธิ์", "SecurityConfig และ REST controllers"),
    ("Service และ policy", "กฎธุรกิจ การตรวจสถานะ และการจัดอันดับ", "CourseServiceImpl, CourseWorkflow, CourseMatcherService"),
    ("Repository และ PostgreSQL", "บันทึกข้อมูลสัมพันธ์และข้อจำกัด", "repositories และ Flyway V1–V4"),
], [3.1, 6.7, 5.0])
figure("3.1", "ภาพรวมองค์ประกอบของระบบ CourseHub", "สร้างภาพจาก doc/diagrams/system-component.mmd และตรวจชื่อเวอร์ชันกับโค้ดก่อนแทรก; วางภาพกว้างเต็มพื้นที่พิมพ์")
heading("3.4 แบบจำลองข้อมูล", 2)
para("ฐานข้อมูลเริ่มต้นมี 12 ตารางหลัก ได้แก่ users, user_profiles, providers, provider_members, platforms, courses, course_prices, categories, course_categories, reviews, saved_courses และ audit_logs การออกแบบใช้ความสัมพันธ์หนึ่งต่อหนึ่งระหว่าง users กับ user_profiles และระหว่าง courses กับ course_prices ความสัมพันธ์หนึ่งต่อหลายระหว่าง providers กับ courses และตารางเชื่อมสำหรับหมวดหมู่ของคอร์ส [5] การกำหนด foreign key, unique, check และ optimistic version ช่วยบังคับความสอดคล้องของข้อมูล ทั้งนี้ข้อจำกัดที่ระดับฐานข้อมูลและกฎที่ระดับบริการต้องพิจารณาร่วมกัน")
figure("3.2", "แผนภาพความสัมพันธ์ข้อมูลหลักของ CourseHub", "แทรก doc/diagrams/er-diagram.png ในหน้ากระดาษแนวนอนหรือแยกเป็นภาพขยาย เพราะภาพต้นฉบับมีรายละเอียดมาก")
heading("3.5 กระบวนการจับคู่คอร์ส", 2)
para("คำขอระบุหมวดหมู่ ระดับ ภาษา งบประมาณเป็นบาท และชั่วโมงเรียนต่อสัปดาห์ ระบบอ่านเฉพาะคอร์สสถานะ PUBLISHED จาก Provider สถานะ ACTIVE จากนั้น EligibilityPolicy ตรวจแต่ละข้อจำกัดอย่างอิสระ คอร์สที่ราคายังไม่ทราบ ใช้สกุลเงินอื่น หรือเป็น subscription ที่ยังคำนวณค่าใช้จ่ายรวมไม่ได้จะไม่เข้าสู่การจัดอันดับ [2,4]")
para("คอร์สที่ผ่านการกรองได้รับคะแนนสามส่วนในช่วง 0–100 ได้แก่ คะแนนราคา B = 100 เมื่อคอร์สฟรี หรือ 100 × (1 − ราคา/งบ) เมื่อเป็นราคาจ่ายครั้งเดียวในงบ คะแนนเวลา E = 100 เมื่อจำนวนชั่วโมงรวมไม่เกิน 4 เท่าของชั่วโมงที่เรียนได้ต่อสัปดาห์ มิฉะนั้นใช้ 100 × ชั่วโมงที่มีใน 4 สัปดาห์/ชั่วโมงรวมของคอร์ส ถ้าไม่ทราบเวลาเรียนใช้ค่า 50 คะแนนรีวิว R = คะแนนเฉลี่ยรีวิวที่เผยแพร่/5 × 100 และใช้ค่า 50 เมื่อยังไม่มีรีวิวที่เผยแพร่ [4]")
para("คะแนนรวมเป็นค่าเฉลี่ยเลขคณิตของ B, E และ R ระบบเรียงคะแนนดิบจากมากไปน้อย ใช้รหัสคอร์สตัดสินกรณีคะแนนเท่ากัน แล้วแสดงไม่เกินสามคอร์สพร้อมคะแนนที่ปัดสองตำแหน่งและเหตุผลแต่ละส่วน การใช้ค่า 50 กับข้อมูลเวลาเรียนหรือรีวิวที่ขาดเป็นสมมติฐานเชิงออกแบบที่ควรทดสอบผลกระทบกับข้อมูลจริง ไม่ใช่ค่าที่พิสูจน์แล้วว่าเหมาะสมที่สุด")
heading("3.6 การควบคุมสิทธิ์และวงจรชีวิตข้อมูล", 2)
para("ระบบใช้ session cookie ร่วมกับ CSRF token สำหรับคำขอที่เปลี่ยนข้อมูล สอดคล้องกับคำแนะนำของ Spring Security สำหรับเว็บที่เบราว์เซอร์ส่ง cookie อัตโนมัติ [15] การตรวจสิทธิ์ที่บริการป้องกันการเข้าถึงข้ามผู้ให้บริการ แม้ส่วนหน้าจะซ่อนปุ่มไว้แล้วก็ตาม")
para("คอร์สเริ่มที่ DRAFT ก่อนส่งตรวจเป็น PENDING ผู้ดูแลสามารถเผยแพร่ ขอแก้ไข ระงับ หรือเก็บเข้าคลังตามสถานะที่อนุญาต การใช้ State pattern รวบรวมกฎการเปลี่ยนสถานะ ขณะที่ audit log เก็บการกระทำสำคัญ และ Observer หลัง transaction commit ใช้เพิ่มตัวนับเหตุการณ์ [4] การแยกตัวนับออกจาก audit log หมายความว่าตัวนับไม่ใช่ประวัติถาวรของระบบ")
heading("3.7 แผนการตรวจสอบ", 2)
para("แผนทดสอบแบ่งเป็นหน่วย การเชื่อมต่อ API และฐานข้อมูล ส่วนหน้า และ end-to-end ในเบราว์เซอร์ เกณฑ์ตรวจรับครอบคลุมสถานะ HTTP การตรวจข้อมูล สิทธิ์ CSRF การกรองแค็ตตาล็อก การจัดอันดับ และวงจรการเผยแพร่ [6] การแยกฐานข้อมูลทดสอบใน E2E ลดผลกระทบจากข้อมูลค้าง และช่วยให้ทดสอบ workflow ของผู้เรียน ผู้ให้บริการ และผู้ดูแลร่วมกัน")
page_break()

heading("บทที่ 4 ผลการพัฒนาและอภิปรายผล")
heading("4.1 ฟังก์ชันที่มีหลักฐานในโปรเจกต์", 2)
table(["ส่วนงาน", "สถานะที่ตรวจพบ", "หลักฐานหลัก"], [
    ("บัญชีและโปรไฟล์", "มี API และหน้าใช้งาน", "README; AuthController; ProfilePage"),
    ("แค็ตตาล็อกและ Bookmark", "มีค้นหา กรอง แบ่งหน้า และบันทึกคอร์ส", "CatalogPage; BookmarksPage; integration tests"),
    ("Provider และ Course", "มี CRUD และส่งตรวจตามสิทธิ์", "ProviderPage; CourseManagementSection; tests"),
    ("Matcher", "มีแบบสอบถามและผลลัพธ์พร้อมเหตุผล", "MatcherPage; CourseMatcherService; tests"),
    ("Admin moderation", "มีการตรวจ Provider, Course และ Review", "AdminPage; moderation services; tests"),
    ("รีวิวผู้เรียน", "มี API แต่ยังไม่มีหน้าเขียนหรือแก้รีวิว", "requirements matrix; E2E report"),
], [3.1, 6.0, 5.7])
figure("4.1", "หน้าค้นหาคอร์สและตัวกรอง", "ถ่ายหน้าจอ /courses จากระบบที่รันจริง โดยให้เห็นช่องค้นหา ตัวกรอง และคอร์สที่เผยแพร่; ปิดข้อมูลบัญชีส่วนตัว")
figure("4.2", "หน้าคำถามและผลลัพธ์ตัวจับคู่คอร์ส", "ถ่ายหน้าจอ /match สองภาพย่อย คือคำถามและผลลัพธ์ที่มีคะแนนพร้อมเหตุผล; ใส่คำตอบตัวอย่างในคำบรรยาย")
figure("4.3", "หน้าจัดการคอร์สและการตรวจอนุมัติ", "ถ่ายหน้าจอ /providers และ /admin หลังสร้างข้อมูลทดสอบ โดยให้เห็นสถานะก่อนและหลังการตรวจ")
heading("4.2 ผลการทดสอบที่ตรวจสอบได้", 2)
para("รายงาน Task 22 ลงวันที่ 7 ตุลาคม พ.ศ. 2569 ระบุการรัน Playwright กับ Chromium และ PostgreSQL แยกฐานข้อมูลสองรอบ รอบละ 10 กรณีผ่านทั้งหมด รวมทั้ง backend 308 กรณีและ frontend 31 กรณีผ่าน [7] ในการจัดทำรายงานนี้ตรวจไฟล์ XML ของ Maven ในเครื่องเพิ่มเติม พบ 308 กรณี failures 0, errors 0, skipped 0 และตรวจ results.json ของ Playwright รอบวันที่ 8 ตุลาคม พบ expected 10, unexpected 0, skipped 0, flaky 0 [8]")
table(["ระดับการทดสอบ", "ผลที่มีหลักฐาน", "ความหมายของผล"], [
    ("Backend", "308 ผ่าน; failures/errors/skipped 0", "ยืนยันกรณีหน่วยและการเชื่อมต่อที่ชุดทดสอบครอบคลุม"),
    ("Frontend", "31 ผ่านตามรายงาน Task 22", "ยืนยันพฤติกรรมคอมโพเนนต์ที่รายงานระบุ"),
    ("End-to-end", "10 ผ่านในรอบล่าสุด; unexpected/skipped/flaky 0", "ยืนยัน flow ใน Chromium และสภาพแวดล้อมทดสอบ"),
], [3.3, 5.1, 6.4])
para("กรณี E2E ครอบคลุมการสมัครและเข้าสู่ระบบ การบันทึกคอร์ส การปฏิเสธสิทธิ์ การตรวจรีวิว การจับคู่แบบมีและไม่มีผลลัพธ์ และการเผยแพร่คอร์สจากผู้ให้บริการถึงผู้เรียน [7] ผลเหล่านี้แสดงการทำงานร่วมกันของส่วนหน้า API และฐานข้อมูลในเงื่อนไขที่ทดสอบ แต่ยังไม่แทนการวัดภาระงาน ความพร้อมใช้งานของระบบสาธารณะ หรือการทดลองกับผู้ใช้จริง")
figure("4.4", "หลักฐานผลการทดสอบและ CI", "เพิ่มภาพสรุป XML/JUnit, Playwright HTML report และหน้า GitHub Actions ของ commit ที่จะส่งจริง; ระบุวันที่และ commit ในคำบรรยาย")
heading("4.3 การเชื่อมแนวคิดกับการลงมือพัฒนา", 2)
para("Strategy pattern อยู่ใน ScoringStrategy และ implementation ด้านงบ เวลา และรีวิว ทำให้เพิ่มวิธีคิดคะแนนได้ผ่าน interface ที่ใช้ร่วมกัน State pattern อยู่ใน CourseWorkflow เพื่อกำหนดคำสั่งที่อนุญาตตามสถานะคอร์ส ส่วน Observer ใช้ event หลัง commit เพื่อบันทึกตัวนับการเปลี่ยนสถานะ [4] เอกสาร SOLID ในโปรเจกต์ให้ตัวอย่างที่ตรวจย้อนกลับถึงคลาสและเทสต์ได้ แต่ไม่ได้พิสูจน์ว่าทุกคลาสในระบบเป็นไปตามทุกหลักการอย่างสมบูรณ์")
heading("4.4 ข้อจำกัดของผลการศึกษา", 2)
para("ข้อจำกัดแรกคือยังไม่มีการประเมินความเกี่ยวข้องของคอร์สโดยผู้เรียนหรือผู้เชี่ยวชาญ จึงสรุปได้เพียงว่ากฎการกรองและคะแนนทำงานตามที่กำหนด ข้อจำกัดที่สองคือ Matcher ตัดคอร์สแบบ subscription และราคาต่างสกุลเงินออกเพราะข้อมูลต้นทุนรวมไม่เพียงพอ ซึ่งอาจลดจำนวนคอร์สที่แนะนำ ข้อจำกัดที่สามคือการทดสอบ E2E ในหลักฐานที่ตรวจใช้ Chromium บนสภาพแวดล้อมทดสอบ จึงไม่ยืนยันผลบนมือถือ เบราว์เซอร์อื่น หรือเมื่อ backend cloud เริ่มทำงานหลังพัก [7]")
para("เอกสารแผนการเผยแพร่ระบุ Vercel, Render และ Neon และมี workflow สำหรับ CI/CD แต่การตั้งค่า URL ส่วนหน้าและผล deploy สาธารณะของรุ่นที่จะส่งยังต้องยืนยันแยกต่างหาก [3,16] ภาพ system-architecture.jpg ใน repository ระบุ Spring Boot 4.0.0 ซึ่งไม่ตรงกับ pom.xml ที่ใช้ 3.5.16 รายงานนี้จึงไม่ใช้ภาพดังกล่าวเป็นหลักฐานสถาปัตยกรรมฉบับสุดท้าย")
page_break()

heading("บทที่ 5 สรุปและข้อเสนอแนะ")
heading("5.1 สรุปผล", 2)
para("CourseHub เป็นระบบต้นแบบสำหรับค้นหาและจับคู่คอร์สเรียนที่รวมข้อมูลจากผู้ให้บริการหลายรายภายใต้กระบวนการตรวจอนุมัติ การออกแบบแยกส่วนหน้า API กฎธุรกิจ และฐานข้อมูล พร้อมกำหนดสิทธิ์ตามบทบาทและสถานะข้อมูล ตัวจับคู่ใช้การกรองเงื่อนไขที่จำเป็นก่อนจัดอันดับและคืนเหตุผลของคะแนน ทำให้พฤติกรรมที่ตรวจสอบได้สอดคล้องกับความต้องการที่ระบุในระบบ")
para("หลักฐานการทดสอบที่ตรวจพบสนับสนุนว่าฟังก์ชันหลักและ flow ระหว่างผู้ใช้สามบทบาททำงานผ่านในสภาพแวดล้อมที่รายงาน อย่างไรก็ตามผลดังกล่าวไม่เพียงพอจะอ้างว่าคำแนะนำเหมาะสมกับผู้เรียนจริง หรือระบบพร้อมใช้งานสาธารณะทุกสภาวะ")
heading("5.2 แนวทางพัฒนาต่อ", 2)
para("ควรเก็บชุดข้อมูลคอร์สที่มีคำตัดสินความเกี่ยวข้องจากผู้เรียนหรือผู้เชี่ยวชาญ แล้วเปรียบเทียบผลจัดอันดับของสูตรปัจจุบันกับวิธีฐานที่กำหนดไว้ล่วงหน้า เช่น การเรียงราคาหรือรีวิว โดยรายงานตัวชี้วัดและขนาดตัวอย่างอย่างครบถ้วน ควรทดลองความเข้าใจของเหตุผลประกอบคำแนะนำและเวลาที่ใช้ตัดสินใจผ่านการทดสอบใช้งานจริง [11,13]")
para("ด้านผลิตภัณฑ์ควรเติมหน้าเขียนและแก้รีวิวสำหรับผู้เรียน ตรวจสอบข้อมูลรอบบิลของคอร์ส subscription รองรับสกุลเงินอื่นด้วยอัตราแลกเปลี่ยนที่มีแหล่งอ้างอิง และทดสอบการแสดงผลบนมือถือกับเบราว์เซอร์หลายชนิด ก่อนส่งงานควรยืนยันสถานะ frontend deployment, GitHub Actions, README และภาพสถาปัตยกรรมให้ตรงกับ commit ที่ส่ง")
page_break()

heading("เอกสารอ้างอิง")
refs = [
    "[1] กลุ่ม 36. (2569). คอร์สดีบอกต่อ — ภาพรวมและแผนงาน MVP [ไฟล์ส่งออกจาก Notion, ปรับแผน 15 กันยายน 2569]. แผน notion.pdf; ข้อความแผนฉบับเต็มที่สอดคล้องกันอยู่ใน GROUP36_CourseRecommend/วันที่ปรับแผน 15 กันยายน 2026.txt.",
    "[2] กลุ่ม 36. (2569). README และซอร์สโค้ด CourseHub [GitHub repository]. https://github.com/poohlikung/GROUP36_CourseRecommend (ตรวจ local checkout branch sorawit_6733800648_01 ที่ commit 2899821 วันที่ 8 ตุลาคม 2569).",
    "[3] กลุ่ม 36. (2569). Requirement Matrix — CourseHub. GROUP36_CourseRecommend/doc/requirements.md.",
    "[4] กลุ่ม 36. (2569). Design Patterns — หลักฐาน implementation; SOLID analysis — Task 21. GROUP36_CourseRecommend/doc/design-patterns.md และ doc/solid-analysis.md.",
    "[5] กลุ่ม 36. (2569). Data Dictionary — CourseHub และ Flyway V1–V4. GROUP36_CourseRecommend/doc/data-dictionary.md; code/backend/src/main/resources/db/migration/.",
    "[6] กลุ่ม 36. (2569). Test Plan — CourseHub. GROUP36_CourseRecommend/doc/test-plan.md.",
    "[7] กลุ่ม 36. (2569). Task 22 — รายงาน End-to-end learner/provider/admin. GROUP36_CourseRecommend/doc/test-reports/task22-e2e.md.",
    "[8] กลุ่ม 36. (2569). ผลรันทดสอบในเครื่อง: code/backend/target/surefire-reports/TEST-*.xml และ test/reports/e2e/coursehub-e2e-cc6ba75c/results.json (ตรวจวันที่ 8 ตุลาคม 2569; ไฟล์ build ไม่ได้ commit).",
    "[9] Adomavicius, G., & Tuzhilin, A. (2005). Toward the next generation of recommender systems: A survey of the state-of-the-art and possible extensions. IEEE Transactions on Knowledge and Data Engineering, 17(6), 734–749. https://doi.org/10.1109/TKDE.2005.99",
    "[10] Burke, R. (2002). Hybrid recommender systems: Survey and experiments. User Modeling and User-Adapted Interaction, 12, 331–370. https://doi.org/10.1023/A:1021240730564",
    "[11] Tintarev, N., & Masthoff, J. (2010). Designing and evaluating explanations for recommender systems. In Recommender Systems Handbook (pp. 479–510). Springer. https://doi.org/10.1007/978-0-387-85820-3_15",
    "[12] International Organization for Standardization. (2023). ISO/IEC 25010:2023 Systems and software engineering — Product quality model. https://committee.iso.org/standard/78176.html",
    "[13] Shani, G., & Gunawardana, A. (2010). Evaluating recommendation systems. In Recommender Systems Handbook (pp. 257–297). Springer. https://doi.org/10.1007/978-0-387-85820-3_8",
    "[14] Fowler, M. (2003). Catalog of Patterns of Enterprise Application Architecture. https://martinfowler.com/eaaCatalog/",
    "[15] Spring Security. (2026). Cross Site Request Forgery (CSRF). https://docs.spring.io/spring-security/reference/6.5/servlet/exploits/csrf.html",
    "[16] กลุ่ม 36. (2569). Continuous deployment (CD). GROUP36_CourseRecommend/doc/deployment-cd.md.",
]
for ref in refs:
    para(ref, "Reference")
page_break()

heading("ภาคผนวก ก รายการภาพและข้อมูลที่ต้องยืนยันก่อนส่ง")
para("ภาคผนวกนี้เป็นรายการสำหรับปรับฉบับร่างให้พร้อมส่ง ไม่ใช่ผลการทดลองเพิ่มเติม")
table(["รายการ", "ตำแหน่ง", "สิ่งที่ต้องทำ"], [
    ("ภาพองค์ประกอบระบบ", "รูปที่ 3.1", "แปลง system-component.mmd เป็นภาพคมชัด ตรวจชื่อเวอร์ชันและองค์ประกอบ"),
    ("ERD", "รูปที่ 3.2", "แทรก er-diagram.png แบบขยายหรือหน้าแนวนอน"),
    ("ภาพหน้าจอ Catalog", "รูปที่ 4.1", "ถ่ายจากระบบจริงพร้อมวันที่และข้อมูลตัวอย่างที่อธิบายได้"),
    ("ภาพ Matcher", "รูปที่ 4.2", "แสดงคำตอบตัวอย่าง คะแนนและเหตุผล; ปิดข้อมูลส่วนบุคคล"),
    ("ภาพ Provider/Admin", "รูปที่ 4.3", "แสดงลำดับสถานะ draft, pending, published หรือ revision"),
    ("ภาพผลทดสอบ", "รูปที่ 4.4", "ใช้ผลของ commit สุดท้าย ไม่ใช้ภาพ 239 tests ที่เก่ากว่า"),
], [3.4, 2.3, 8.1])
para("ข้อมูลหน้าปกที่ต้องยืนยัน: ชื่อรายวิชาและรหัสวิชา ภาคเรียน ปีการศึกษา ชื่ออาจารย์ผู้สอนหรือที่ปรึกษา วันที่ส่ง และรูปแบบชื่อสมาชิกตามทะเบียน")
para("ข้อมูลโครงการที่ต้องยืนยัน: URL frontend สาธารณะ, สถานะ CI/CD ของ commit ที่ส่ง, branch/commit สุดท้าย, ผลทดสอบที่รันซ้ำหลังแก้ไขล่าสุด และนโยบายการใช้เครื่องมือ AI ของรายวิชา โดยเฉพาะข้อความเปิดเผยการใช้เครื่องมือในงานส่ง")
para("คำเตือนด้านภาพ: ภาพ system-architecture.jpg ใน repository ระบุ Spring Boot 4.0.0 และภาพ use-case-diagram.png รวม Roadmap ซึ่งยังไม่ใช่ฟังก์ชันที่พัฒนาเสร็จ จึงควรแก้ภาพก่อนใช้เป็นหลักฐานของระบบปัจจุบัน")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Cm(0)
font(footer.add_run("CourseHub  |  "), 11, color=GRAY)
f = OxmlElement("w:fldSimple")
f.set(qn("w:instr"), "PAGE")
footer._p.append(f)

doc.core_properties.title = "รายงานวิชาการ การออกแบบและพัฒนาระบบแนะนำคอร์สเรียนคอร์สดีบอกต่อ"
doc.core_properties.subject = "CourseHub academic project report"
doc.core_properties.author = "กลุ่ม 36"
doc.core_properties.keywords = "CourseHub; course recommendation; software design"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
